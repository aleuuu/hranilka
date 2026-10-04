"""Движок Хранилки: сканирование, понимание файлов локальной моделью, группировка, план, дотюнивание.

Протокол — JSON-строки через stdin/stdout:
  запрос  {"id": 1, "cmd": "scan", ...}
  ответ   {"id": 1, "ok": true, "result": ...} | {"id": 1, "ok": false, "error": "..."}
  событие {"event": "progress", ...}
Файлы движок только читает. Перемещением занимается приложение (Rust), с журналом отмены.
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import hashlib
import html as htmllib
import io
import json
import os
import re
import sys
import threading
import time
import traceback
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from pathlib import Path

# OpenBLAS (numpy) на Windows может зависнуть, если впервые грузится не из главного потока
# и успевает создать свои потоки внутри загрузчика DLL. Поэтому: один поток BLAS и импорт здесь, в главном потоке.
os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")
os.environ.setdefault("OMP_NUM_THREADS", "1")

import numpy as np
import requests
from PIL import Image, ImageOps
from pillow_heif import register_heif_opener

register_heif_opener()
Image.MAX_IMAGE_PIXELS = 200_000_000

OLLAMA = os.environ.get("HRANILKA_OLLAMA", "http://127.0.0.1:11434").rstrip("/")
MODELS = {"fast": "gemma3:4b", "accurate": "gemma3:12b"}
EMBED = "bge-m3"
CTX = 8192  # один размер контекста на все вызовы: смена num_ctx заставляет Ollama перезагружать модель
WORKERS = 2
CACHE_VER = 2  # меняется вместе с промптами описания — старые описания тогда пересчитываются

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
DOC_EXT = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".rtf", ".html", ".htm", ".eml", ".csv", ".pptx", ".odt"}
VIDEO_EXT = {".mp4", ".mov", ".webm", ".mkv", ".avi", ".m4v", ".3gp"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac", ".opus"}
INSTALLER_EXT = {".exe", ".msi", ".dmg", ".iso", ".msix", ".appx", ".apk"}
ARCHIVE_EXT = {".zip", ".rar", ".7z", ".tar", ".gz", ".tgz", ".bz2", ".xz"}
TEMP_EXT = {".crdownload", ".part", ".tmp", ".partial", ".download", ".загрузка", ".opdownload"}
CODE_EXT = {".php", ".js", ".ts", ".py", ".rs", ".cs", ".java", ".go", ".sql", ".jsx", ".tsx", ".vue", ".kt", ".swift", ".cpp", ".c", ".h"}
SERVICE_EXT = CODE_EXT | {".css", ".json", ".yaml", ".yml", ".toml", ".xml", ".bat", ".ps1", ".sh", ".cmd", ".conf", ".ini", ".cfg",
                          ".winmd", ".ico", ".svg", ".dll", ".sys", ".log", ".lnk", ".url", ".torrent", ".reg", ".map"}
BINARY_EXT = {".exe", ".dll", ".msi", ".sys"}

DUMP_DIRS = {"telegram desktop", "whatsapp", "screenshots", "снимки экрана", "camera", "dcim", "camera roll",
             "загрузки", "downloads", "телеграм", "photos", "фото"}
PROJECT_MARKERS = {".git", "package.json", "cargo.toml", "composer.json", "pyproject.toml", "requirements.txt",
                   "pom.xml", "build.gradle", "index.php", "go.mod", "makefile", ".vscode", ".idea", "node_modules",
                   "venv", ".venv", "gradlew", "cmakelists.txt"}
SKIP_NAMES = {"desktop.ini", "thumbs.db", ".ds_store"}

KNOWN = {"downloads": "Загрузки", "desktop": "Рабочий стол", "pictures": "Изображения", "documents": "Документы",
         "videos": "Видео", "music": "Музыка"}

COL = {"v": "oklch(0.74 0.15 300)", "p": "oklch(0.74 0.15 350)", "b": "oklch(0.74 0.15 255)",
       "t": "oklch(0.78 0.12 185)", "y": "oklch(0.84 0.13 85)"}

MON = ["янв", "фев", "мар", "апр", "мая", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]

OUT_LOCK = threading.Lock()


def emit(obj: dict) -> None:
    line = json.dumps(obj, ensure_ascii=False)
    with OUT_LOCK:
        sys.stdout.write(line + "\n")
        sys.stdout.flush()


def plural(n: int, forms: tuple[str, str, str]) -> str:
    a = abs(n) % 100
    b = a % 10
    if 10 < a < 20:
        return forms[2]
    if 1 < b < 5:
        return forms[1]
    if b == 1:
        return forms[0]
    return forms[2]


def display_name(p: Path) -> str:
    return KNOWN.get(p.name.lower(), p.name) if p.parent == Path.home() else p.name


def fmt_date(iso: str) -> str:
    try:
        d = dt.date.fromisoformat(iso)
        return f"{d.day} {MON[d.month - 1]} {d.year}"
    except Exception:
        return "—"


# ───────────────────────────── сканирование ─────────────────────────────

@dataclass
class Item:
    id: int
    path: Path
    rel: str                 # путь относительно корня через '/'
    kind: str                # image | doc | video | audio | rule
    size: int
    mtime: float
    ext: str
    rule: str = ""           # для файлов из «Вернуть в разбор»: installers | archives | temp | service | other
    meta: dict = field(default_factory=dict)
    ai: dict = field(default_factory=dict)
    date: str = ""
    date_src: str = ""
    thumb: str | None = None
    category: str = ""
    cluster: str = ""
    vec: object = None       # смысловой отпечаток (для обучения на правках)

    @property
    def from_rel(self) -> str:
        parent = str(Path(self.rel).parent).replace("\\", "/")
        return "" if parent == "." else parent


def classify_folder(p: Path) -> str:
    """Категория неделимой папки или '' — если её можно разбирать по файлам."""
    if p.name.lower() in DUMP_DIRS:
        return ""
    names, exts, n = set(), Counter(), 0
    try:
        for f in p.rglob("*"):
            names.add(f.name.lower())
            if f.is_file():
                exts[f.suffix.lower()] += 1
                n += 1
            if n > 4000:
                break
    except OSError:
        pass
    if names & PROJECT_MARKERS or any(nm.endswith(".sln") or nm.endswith(".csproj") for nm in names):
        return "code"
    if sum(exts[e] for e in CODE_EXT) >= 3:
        return "code"
    if sum(exts[e] for e in BINARY_EXT) >= 2:
        return "programs"
    return "userdirs"


def file_category(e: str) -> str:
    if e in IMAGE_EXT: return "image"
    if e in DOC_EXT: return "doc"
    if e in VIDEO_EXT: return "video"
    if e in AUDIO_EXT: return "audio"
    if e in INSTALLER_EXT: return "installers"
    if e in ARCHIVE_EXT: return "archives"
    if e in TEMP_EXT: return "temp"
    if e in SERVICE_EXT: return "service"
    return "other"


FOLDER_WHY = {"code": "код", "programs": "программа", "userdirs": "ваша папка"}
CATS = {
    "code": ("code", "Проекты с кодом"), "programs": ("apps", "Программы"), "userdirs": ("folder_special", "Ваши папки"),
    "webpages": ("language", "Сохранённые веб-страницы"), "installers": ("install_desktop", "Установщики"),
    "archives": ("folder_zip", "Архивы"), "temp": ("hourglass_empty", "Временные файлы"),
    "service": ("settings", "Служебные файлы"), "other": ("help", "Другие типы файлов"),
    "excluded": ("block", "Исключено в настройках"),
}


@dataclass
class Scan:
    root: Path
    items: list[Item]
    locked: list[dict]
    skipped: dict[str, list[str]]  # категория -> примеры (имена)
    skipped_counts: Counter
    root_files: int
    root_kinds: list[str]


def code_why(p: Path) -> str:
    exts = Counter()
    try:
        for f in p.rglob("*"):
            if f.is_file():
                exts[f.suffix.lower()] += 1
            if sum(exts.values()) > 2000:
                break
    except OSError:
        pass
    langs = {".php": "PHP", ".py": "Python", ".js": "JavaScript", ".ts": "TypeScript", ".rs": "Rust", ".cs": "C#",
             ".java": "Java", ".go": "Go", ".kt": "Kotlin", ".swift": "Swift", ".cpp": "C++", ".jsx": "React", ".tsx": "React"}
    best = max(((exts[e], name) for e, name in langs.items()), default=(0, ""))
    return f"код · {best[1]}" if best[0] else "код"


def count_files(p: Path, limit=100000) -> int:
    n = 0
    try:
        for _, _, files in os.walk(p):
            n += len(files)
            if n > limit:
                break
    except OSError:
        pass
    return n


def scan(root: Path, unskip: set[str], exclude: set[str] | None = None) -> Scan:
    exclude = {e.lower().rstrip("\\/") for e in (exclude or set())}
    if not root.exists() or not root.is_dir():
        raise RuntimeError("Папка не найдена или недоступна")
    parts = [p.lower() for p in root.resolve().parts]
    if len(parts) >= 2 and parts[1] in {"windows", "program files", "program files (x86)", "programdata"}:
        raise RuntimeError("Это системная папка — её разбирать нельзя")
    r = str(root).lower().rstrip("\\/")
    if any(r == e or r.startswith(e + "\\") for e in exclude):
        raise RuntimeError("Эта папка в списке «Не трогать никогда» — уберите её оттуда в настройках, если хотите разобрать")
    items: list[Item] = []
    locked: list[dict] = []
    skipped: dict[str, list[str]] = defaultdict(list)
    counts: Counter = Counter()
    nid = 0
    root_files = 0
    root_kinds: list[str] = []
    html_with_files: set[str] = set()

    def add(p: Path, kind: str, rule: str = ""):
        nonlocal nid
        try:
            st = p.stat()
        except OSError:
            return
        nid += 1
        items.append(Item(id=nid, path=p, rel=str(p.relative_to(root)).replace("\\", "/"), kind=kind, size=st.st_size,
                          mtime=st.st_mtime, ext=p.suffix.lower(), rule=rule))

    def walk(d: Path, depth: int):
        nonlocal root_files
        try:
            entries = sorted(d.iterdir(), key=lambda x: x.name.lower())
        except OSError:
            return
        # пары «страница.html + страница_files»
        dirs = {e.name for e in entries if e.is_dir()}
        for e in entries:
            if e.is_file() and e.suffix.lower() in {".html", ".htm"} and (e.stem + "_files") in dirs:
                html_with_files.add(str(e))
        for e in entries:
            if e.name.lower() in SKIP_NAMES or e.name.startswith("~$"):
                continue
            if str(e).lower().rstrip("\\/") in exclude:
                skipped["excluded"].append(e.name)
                counts["excluded"] += count_files(e) if e.is_dir() else 1
                continue
            if e.is_dir():
                if e.name.endswith("_files") and any(str(d / (e.name[:-6] + x)) in html_with_files for x in (".html", ".htm")):
                    if "webpages" in unskip:
                        walk(e, depth + 1)
                    else:
                        skipped["webpages"].append(e.name[:-6]); counts["webpages"] += count_files(e)
                    continue
                cat = classify_folder(e)
                if not cat:
                    walk(e, depth + 1)
                elif cat in unskip:
                    walk(e, depth + 1)
                else:
                    why = code_why(e) if cat == "code" else FOLDER_WHY[cat]
                    locked.append({"id": f"L{len(locked) + 1}", "name": e.name, "why": why, "rel": str(e.relative_to(root)).replace("\\", "/"),
                                   "abs": str(e), "cur": str(e), "n": count_files(e), "cat": cat})
                    skipped[cat].append(e.name)
                continue
            if depth == 0:
                root_files += 1
            cat = file_category(e.suffix.lower())
            if str(e) in html_with_files and "webpages" not in unskip:
                counts["webpages"] += 1
                continue
            if depth == 0 and len(root_kinds) < 84:
                root_kinds.append({"image": "image", "doc": "doc", "video": "video", "audio": "video"}.get(cat, "manual"))
            if cat in {"image", "doc", "video", "audio"}:
                add(e, cat)
            elif cat in unskip:
                add(e, "rule", cat)
            else:
                skipped[cat].append(e.name); counts[cat] += 1

    walk(root, 0)
    return Scan(root=root, items=items, locked=locked, skipped=skipped, skipped_counts=counts, root_files=root_files, root_kinds=root_kinds)


def scan_summary(sc: Scan) -> dict:
    kinds = Counter(i.kind for i in sc.items)
    cats = []
    for key, (icon, label) in CATS.items():
        names = sc.skipped.get(key, [])
        if key in {"code", "programs", "userdirs"}:
            if not names:
                continue
            ex = ", ".join(names[:5]) + (f" и ещё {len(names) - 5}" if len(names) > 5 else "")
            cats.append({"key": key, "icon": icon, "label": label, "ex": ex, "count": len(names)})
        else:
            n = sc.skipped_counts.get(key, 0)
            if not n:
                continue
            if key in ("webpages", "excluded"):
                ex = ", ".join(names[:4]) or f"{n} файлов"
            else:
                exts = Counter(Path(x).suffix.lower() for x in names)
                ex = f"{n} {plural(n, ('файл', 'файла', 'файлов'))}: " + ", ".join(e for e, _ in exts.most_common(4) if e)
            cats.append({"key": key, "icon": icon, "label": label, "ex": ex, "count": n})
    skipped_files = sum(v for k, v in sc.skipped_counts.items())
    return {
        "root": str(sc.root), "rootDisplay": display_name(sc.root), "total": len(sc.items),
        "images": kinds["image"], "docs": kinds["doc"], "media": kinds["video"] + kinds["audio"],
        "lockedCount": len(sc.locked), "skippedFiles": skipped_files, "skipped": cats,
    }


# ───────────────────────────── метаданные и превью ─────────────────────────────

DATE_IN_NAME = [
    re.compile(r"(20\d{2})[-_.]?(0[1-9]|1[0-2])[-_.]?(0[1-9]|[12]\d|3[01])"),
    re.compile(r"(0[1-9]|[12]\d|3[01])\.(0[1-9]|1[0-2])\.(20\d{2})"),
]


def date_from_name(name: str) -> str:
    for i, rx in enumerate(DATE_IN_NAME):
        m = rx.search(name)
        if m:
            y, mo, d = (m.group(1), m.group(2), m.group(3)) if i == 0 else (m.group(3), m.group(2), m.group(1))
            if 2000 <= int(y) <= dt.date.today().year:
                return f"{y}-{mo}-{d}"
    return ""


def open_image(it: Item) -> Image.Image | None:
    try:
        img = Image.open(it.path)
        exif = img.getexif()
        sub = exif.get_ifd(0x8769)
        dto = sub.get(36867) or exif.get(306)
        parts = [str(x).replace("\x00", "").strip() for x in (exif.get(271), exif.get(272)) if x]
        cam = parts[-1] if len(parts) == 2 and parts[1].lower().startswith(parts[0].lower()) else " ".join(parts)
        it.meta.update(width=img.width, height=img.height, camera=cam)
        if dto and re.match(r"\d{4}:\d{2}:\d{2}", str(dto)):
            d = str(dto)[:10].replace(":", "-")
            if d.endswith("-01-01") and int(d[:4]) <= 2015:
                it.meta["exif_reset"] = d
            elif 1990 < int(d[:4]) <= dt.date.today().year:
                it.date, it.date_src = d, "EXIF"
        return ImageOps.exif_transpose(img)
    except Exception as e:
        it.meta["error"] = str(e)[:160]
        return None


def to_b64(img: Image.Image, side: int, fmt: str, q: int) -> str:
    im = img.copy()
    if getattr(im, "n_frames", 1) > 1:
        im.seek(0)
    im = im.convert("RGB")
    im.thumbnail((side, side))
    buf = io.BytesIO()
    im.save(buf, fmt, quality=q)
    return base64.b64encode(buf.getvalue()).decode()


THUMBS: Path | None = None  # папка для превью; задаётся в main()


def thumb_url(img: Image.Image, key: str = "") -> str:
    """Сохраняет превью файлом в кэш приложения и возвращает путь (интерфейс показывает его через asset-протокол).
    Без папки кэша — встраивает картинку прямо в ответ."""
    data = to_b64(img, 320, "WEBP", 70)
    if THUMBS is None:
        return "data:image/webp;base64," + data
    THUMBS.mkdir(parents=True, exist_ok=True)
    name = hashlib.sha1((key or data[:2000]).encode("utf-8", "ignore")).hexdigest() + ".webp"
    f = THUMBS / name
    if not f.exists():
        f.write_bytes(base64.b64decode(data))
    return str(f)


def doc_text(it: Item, limit: int = 3000) -> str:
    e, p = it.ext, it.path
    try:
        if e == ".pdf":
            import pymupdf
            with pymupdf.open(p) as d:
                it.meta["pages"] = d.page_count
                text = "".join(d[i].get_text() for i in range(min(3, d.page_count)))
                try:
                    pix = d[0].get_pixmap(dpi=40)
                    im = Image.open(io.BytesIO(pix.tobytes("png")))
                    it.thumb = thumb_url(im, f"{it.path}|{it.size}|{int(it.mtime)}|pdf")
                    if not text.strip():
                        it.meta["scan_b64"] = to_b64(Image.open(io.BytesIO(d[0].get_pixmap(dpi=110).tobytes("png"))), 1024, "JPEG", 85)
                except Exception:
                    pass
                return text[:limit]
        if e == ".docx":
            import docx
            d = docx.Document(str(p))
            return "\n".join(x.text for x in d.paragraphs if x.text.strip())[:limit]
        if e == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:3]:
                out.append(f"[Лист {ws.title}]")
                for row in ws.iter_rows(max_row=15, values_only=True):
                    out.append(" | ".join(str(c) for c in row if c is not None))
            return "\n".join(out)[:limit]
        if e in {".html", ".htm"}:
            raw = p.read_text("utf-8", errors="ignore")
            if "__bundler/manifest" in raw[:200000] or "__bundler_loading" in raw[:20000]:
                # упакованная выгрузка дизайн-макета: внутри только загрузчик, смысл — в имени файла
                return (f"Интерактивный прототип интерфейса (дизайн-макет), сохранённый одним HTML-файлом. "
                        f"Содержимое упаковано, поэтому название и проект бери из имени файла «{p.stem}».")
            title = re.search(r"<title[^>]*>(.*?)</title>", raw, re.S | re.I)
            body = re.sub(r"<(script|style|noscript)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
            body = re.sub(r"<[^>]+>", " ", body)
            body = htmllib.unescape(re.sub(r"\s+", " ", body))
            return (f"Заголовок: {title.group(1).strip()}\n" if title else "") + body[:limit]
        if e in {".pptx", ".odt", ".rtf"}:
            return ""
        return p.read_text("utf-8", errors="ignore")[:limit]
    except Exception as ex:
        it.meta["error"] = str(ex)[:160]
        return ""


MOJIBAKE = re.compile(r"[ÐÑ][\x80-\xbf\u0080-ÿА-я]")


def name_is_broken(name: str) -> bool:
    return bool(MOJIBAKE.search(name)) or "Ð" in name


# ───────────────────────────── модель ─────────────────────────────

def ollama_chat(model: str, prompt: str, schema: dict | None, images: list[str] | None = None, retries: int = 2) -> tuple[dict, dict]:
    msg = {"role": "user", "content": prompt}
    if images:
        msg["images"] = images
    body = {"model": model, "messages": [msg], "stream": False, "options": {"temperature": 0.2, "num_ctx": CTX}, "keep_alive": "30m"}
    if schema:
        body["format"] = schema
    last = None
    for attempt in range(retries + 1):
        try:
            r = requests.post(f"{OLLAMA}/api/chat", json=body, timeout=900)
            if r.status_code == 404:
                raise RuntimeError(f"Модель {model} не установлена")
            r.raise_for_status()
            j = r.json()
            content = j["message"]["content"]
            return (json.loads(content) if schema else {"text": content}), j
        except requests.ConnectionError:
            last = RuntimeError("Движок Ollama не отвечает — он запущен?")
            time.sleep(1.5)
        except Exception as e:
            last = e
            time.sleep(0.5 * (attempt + 1))
    raise RuntimeError(str(last))


def embed(texts: list[str], batch: int = 32) -> np.ndarray:
    out = []
    for i in range(0, len(texts), batch):
        r = requests.post(f"{OLLAMA}/api/embed", json={"model": EMBED, "input": texts[i:i + batch], "keep_alive": "30m"}, timeout=600)
        r.raise_for_status()
        out.extend(r.json()["embeddings"])
    v = np.array(out, dtype=np.float32)
    n = np.linalg.norm(v, axis=1, keepdims=True)
    n[n == 0] = 1
    return v / n


IMG_SCHEMA = {
    "type": "object",
    "properties": {
        "description": {"type": "string"},
        "kind": {"type": "string", "enum": ["фото", "скриншот", "скан документа", "дизайн/макет", "картинка/мем", "схема/график", "другое"]},
        "text_on_image": {"type": "string"},
        "topic": {"type": "string"},
        "suggested_name": {"type": "string"},
        "date_on_image": {"type": "string"},
        "sensitive": {"type": "boolean"},
        "confidence": {"type": "integer"},
    },
    "required": ["description", "kind", "text_on_image", "topic", "suggested_name", "date_on_image", "sensitive", "confidence"],
}

DOC_SCHEMA = {
    "type": "object",
    "properties": {
        "description": {"type": "string"},
        "kind": {"type": "string", "enum": ["другое", "заметка", "прототип/макет", "техзадание/спецификация", "учёба", "отчёт/презентация",
                                             "инструкция/справка", "счёт/чек/финансы", "договор/юридическое", "резюме", "переписка",
                                             "личный документ", "медиа"]},
        "topic": {"type": "string"},
        "suggested_name": {"type": "string"},
        "doc_date": {"type": "string"},
        "sensitive": {"type": "boolean"},
        "confidence": {"type": "integer"},
    },
    "required": ["description", "kind", "topic", "suggested_name", "doc_date", "sensitive", "confidence"],
}


def lang_line(lang: str) -> str:
    return "Пиши description и topic по-русски. suggested_name — по-английски." if lang == "en" else "Отвечай по-русски."


IMG_PROMPT = """Ты помогаешь навести порядок в личных файлах. Посмотри на изображение и ответь JSON. {lang}
Контекст файла: имя «{name}», папка «{folder}», размер {w}x{h}{cam}.
- description: одно конкретное предложение, что изображено (какое приложение или сайт, что за документ, место, люди, предметы).
- kind: что это за изображение.
- text_on_image: ключевой текст на изображении, до 15 слов, или пустая строка.
- topic: тема в 2-4 словах, общая для похожих файлов. Если это интерфейс конкретного приложения, сайта или бренда — назови его прямо («приложение Klutz», «сайт Termoland»).
- suggested_name: понятное имя файла, 2-7 слов, без расширения. Если на картинке есть название приложения/компании — начни с него («Klutz — экран оплаты»). Для чеков и оплат добавь месяц и год в конце («Оплата Термоленд — авг 2026»).
- date_on_image: дата, если она явно видна на изображении (чек, документ), в формате YYYY-MM-DD, иначе пустая строка.
- sensitive: true ТОЛЬКО если это фото или скан документа, удостоверяющего личность (паспорт, ID-карта, СНИЛС, водительские права, полис, свидетельство), банковской карты с номером или медицинского документа. Обычные фото людей, вечеринок, скриншоты приложений, переписок и банковских операций — false.
- confidence: насколько ты уверен в описании и имени, от 0 до 100."""

DOC_PROMPT = """Ты помогаешь навести порядок в личных файлах. По имени и фрагменту содержимого определи документ и ответь JSON. {lang}
Имя файла: «{name}», папка «{folder}».
Фрагмент содержимого:
<<<
{text}
>>>
- description: одно предложение, о чём документ.
- kind: тип документа. «договор/юридическое» — только для настоящих договоров и юридических документов.
- topic: тема в 2-4 словах, общая для похожих файлов. Если документ относится к конкретному проекту, приложению или компании — назови их прямо («проект Termoland»).
- suggested_name: понятное имя файла, 2-8 слов, без расширения. Для счетов, договоров, чеков добавь месяц и год в конце («Договор аренды — сен 2026»). Если имя файла испорчено (кракозябры) — восстанови по содержимому.
- doc_date: дата документа в формате YYYY-MM-DD, если она явно указана в тексте, иначе пустая строка.
- sensitive: true ТОЛЬКО для документов, удостоверяющих личность, и медицинских документов (паспорт, СНИЛС, полис, справка о здоровье). Чеки, счета, договоры, учебные работы — false.
- confidence: насколько ты уверен, от 0 до 100."""

MEDIA_PROMPT = """Ты помогаешь навести порядок в личных файлах. Это {what} — содержимое не анализируется, есть только имя файла «{name}» и папка «{folder}». Ответь JSON. {lang}
- description: одно предложение, что это скорее всего за файл (по имени).
- kind: «медиа».
- topic: тема в 2-4 словах.
- suggested_name: понятное имя, 2-7 слов, без расширения и без даты.
- doc_date: пустая строка.
- sensitive: false.
- confidence: насколько ты уверен, от 0 до 100."""


# ───────────────────────────── состояние ─────────────────────────────

class Cache:
    def __init__(self, path: Path):
        self.path = path
        self.lock = threading.Lock()
        self.data: dict = {}
        self.dirty = 0
        try:
            self.data = json.loads(path.read_text("utf-8"))
        except Exception:
            self.data = {}

    def key(self, it: Item, model: str, lang: str) -> str:
        return f"v{CACHE_VER}|{it.path}|{it.size}|{int(it.mtime)}|{model}|{lang}"

    def get(self, k: str):
        with self.lock:
            return self.data.get(k)

    def put(self, k: str, v: dict):
        with self.lock:
            self.data[k] = v
            self.dirty += 1
            if self.dirty >= 20:
                self._save()

    def save(self):
        with self.lock:
            self._save()

    def _save(self):
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            tmp = self.path.with_suffix(".tmp")
            tmp.write_text(json.dumps(self.data, ensure_ascii=False), "utf-8")
            tmp.replace(self.path)
            self.dirty = 0
        except Exception:
            pass


class Engine:
    def __init__(self, data_dir: Path):
        self.data_dir = data_dir
        self.cache = Cache(data_dir / "cache" / "describe.json")
        self.paused = threading.Event()
        self.stop = threading.Event()
        self.job: threading.Thread | None = None
        self.ctx: dict | None = None  # последний анализ: items, clusters, scan, options
        self.learned_path = data_dir / "learned.json"
        self.learned: list[dict] = []
        try:
            for e in json.loads(self.learned_path.read_text("utf-8")):
                self.learned.append({**e, "v": np.array(e["v"], dtype=np.float32)})
        except Exception:
            self.learned = []

    # ─── обучение на правках ───
    def save_learned(self):
        try:
            self.learned_path.parent.mkdir(parents=True, exist_ok=True)
            data = [{k: (v.tolist() if k == "v" else v) for k, v in e.items()} for e in self.learned]
            self.learned_path.write_text(json.dumps(data, ensure_ascii=False), "utf-8")
        except Exception:
            traceback.print_exc(file=sys.stderr)

    def learn(self, req: dict) -> dict:
        """Пользователь перетащил файлы в папку и нажал «Запомнить»: запоминаем их отпечатки и папку."""
        if not self.ctx:
            raise RuntimeError("План устарел — запустите анализ папки заново")
        folder = str(req.get("folder", "")).strip().strip("/")
        ids = set(int(i) for i in req.get("ids") or [])
        items = [it for it in self.ctx["items"] if it.id in ids]
        need = [it for it in items if it.vec is None]
        if need:
            arr = embed([embed_text(it) for it in need])
            for it, v in zip(need, arr):
                it.vec = v
        for it in items:
            self.learned = [e for e in self.learned if e.get("label") != it.path.name]  # повторная правка — переучиваем
            self.learned.append({"v": np.asarray(it.vec, dtype=np.float32), "folder": folder, "label": it.path.name,
                                 "at": dt.date.today().isoformat()})
        self.learned = self.learned[-500:]
        self.save_learned()
        self.ctx["learned"] = self.learned
        return {"count": len(self.learned)}

    def learned_info(self) -> dict:
        folders = Counter(e["folder"] for e in self.learned)
        return {"count": len(self.learned), "folders": [{"folder": f, "n": n} for f, n in folders.most_common(8)]}

    def learned_reset(self) -> dict:
        self.learned = []
        self.save_learned()
        if self.ctx:
            self.ctx["learned"] = []
        return {"count": 0}

    def replan(self, req: dict) -> dict:
        """Правила или проекты поменялись — пересобираем текущий план без модели."""
        if not self.ctx:
            raise RuntimeError("План устарел — запустите анализ папки заново")
        ctx = self.ctx
        ctx["rules"] = req.get("rules") or []
        ctx["projects_user"] = req.get("projects") or []
        mapping = dict(req.get("mapping") or ctx["mapping"])
        naming = dict(req.get("naming") or ctx["naming"])
        plan = build_plan(ctx, mapping, naming)
        return {"files": plan["files"], "projects": plan["projects"]}

    # ─── описание одного файла ───
    def describe(self, it: Item, model: str, lang: str) -> None:
        folder = it.from_rel or "корень"
        if it.kind == "rule":
            it.ai = {"description": RULE_DESC[it.rule], "kind": "другое", "topic": RULE_DIR[it.rule], "suggested_name": it.path.stem,
                     "sensitive": False, "confidence": 90, "text_on_image": "", "date_on_image": "", "doc_date": ""}
        elif it.kind == "image":
            img = open_image(it)
            if img is None:
                raise RuntimeError("не удалось открыть изображение")
            it.thumb = thumb_url(img, f"{it.path}|{it.size}|{int(it.mtime)}")
            cam = f", камера {it.meta['camera']}" if it.meta.get("camera") else ""
            prompt = IMG_PROMPT.format(name=it.path.name, folder=folder, w=img.width, h=img.height, cam=cam, lang=lang_line(lang))
            it.ai, _ = ollama_chat(model, prompt, IMG_SCHEMA, images=[to_b64(img, 896, "JPEG", 85)])
            d = it.ai.get("date_on_image") or ""
            if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", d) and not it.date:
                it.date, it.date_src = d, "текст на картинке"
        elif it.kind == "doc":
            text = doc_text(it)
            if text.strip():
                it.meta["text"] = text[:1500]  # для поиска ключевых слов проектов и правил
            scan_b64 = it.meta.pop("scan_b64", None)
            if scan_b64:  # PDF-скан без текста — смотрим на первую страницу глазами
                prompt = IMG_PROMPT.format(name=it.path.name, folder=folder, w=0, h=0, cam="", lang=lang_line(lang))
                ai, _ = ollama_chat(model, prompt, IMG_SCHEMA, images=[scan_b64])
                it.ai = {**ai, "kind": "личный документ" if ai.get("sensitive") else "другое", "doc_date": ai.get("date_on_image", "")}
            else:
                if not text.strip():
                    text = "(содержимое недоступно — опирайся на имя файла)"
                it.ai, _ = ollama_chat(model, DOC_PROMPT.format(name=it.path.name, folder=folder, text=text, lang=lang_line(lang)), DOC_SCHEMA)
            d = it.ai.get("doc_date") or ""
            if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", d) and not it.date:
                it.date, it.date_src = d, "текст документа"
        else:
            what = "видео" if it.kind == "video" else "аудиофайл"
            it.ai, _ = ollama_chat(model, MEDIA_PROMPT.format(what=what, name=it.path.name, folder=folder, lang=lang_line(lang)), DOC_SCHEMA)
        if not it.date:
            d = date_from_name(it.path.name)
            if d:
                it.date, it.date_src = d, "имя файла"
        if not it.date:
            it.date = dt.date.fromtimestamp(it.mtime).isoformat()
            it.date_src = "дата изменения"

    # ─── анализ ───
    def analyze(self, req: dict) -> None:
        root = Path(req["root"])
        dest = Path(req.get("dest") or req["root"])
        model = MODELS.get(req.get("model", "accurate"), MODELS["accurate"])
        naming = req.get("naming") or {"rename": True, "datePrefix": True, "lang": "ru", "maxWords": 0}
        wishes = (req.get("wishes") or "").strip()
        unskip = set(req.get("unskip") or [])
        exclude = set(req.get("exclude") or [])
        self.paused.clear()
        self.stop.clear()
        try:
            sc = scan(root, unskip, exclude)
            items = sc.items
            emit({"event": "stage", "stage": 1})
            total = len(items)
            done, errors = 0, 0
            t0 = time.time()
            times: list[float] = []
            topics: Counter = Counter()
            topic_kind: dict[str, Counter] = defaultdict(Counter)
            lang = naming.get("lang", "ru")

            def finish(it: Item, ok: bool):
                nonlocal done, errors
                done += 1
                if not ok:
                    errors += 1
                tp = (it.ai.get("topic") or "").split(",")[0].strip()
                if tp and ok:
                    key = (tp[0].upper() + tp[1:]) if "а" <= tp[0] <= "я" or tp[0] == "ё" else tp
                    topics[key.lower()] += 1
                    topic_kind[key.lower()][type_of(it)[0]] += 1
                    topic_kind[key.lower()]["__label__" + key] += 1
                emit({"event": "file_done", "file": ui_file(it, naming)})
                spf = (time.time() - t0) / max(1, done - cached)
                emit({"event": "progress", "done": done, "total": total, "errors": errors, "secPerFile": round(spf, 2) if done > cached else 0})
                if done % 4 == 0 or done == total:
                    emit({"event": "themes", "themes": themes_of(topics, topic_kind)})

            # из кэша — сразу
            todo = []
            cached = 0
            for it in items:
                c = self.cache.get(self.cache.key(it, model, lang))
                if c:
                    it.ai, it.meta, it.date, it.date_src, it.thumb = c["ai"], c["meta"], c["date"], c["date_src"], c.get("thumb")
                    if it.thumb and not it.thumb.startswith("data:") and not Path(it.thumb).exists():
                        it.thumb = None
                    cached += 1
                    finish(it, True)
                else:
                    todo.append(it)

            def work(it: Item):
                while self.paused.is_set() and not self.stop.is_set():
                    time.sleep(0.3)
                if self.stop.is_set():
                    return it, None
                emit({"event": "file_start", "old": it.path.name, "thumb": None, "kind": type_of(it)[0]})
                s = time.time()
                try:
                    self.describe(it, model, lang)
                    self.cache.put(self.cache.key(it, model, lang), {"ai": it.ai, "meta": it.meta, "date": it.date, "date_src": it.date_src, "thumb": it.thumb})
                    times.append(time.time() - s)
                    return it, True
                except Exception as e:
                    it.ai = {"description": f"Не удалось прочитать файл: {e}"[:200], "kind": "другое", "topic": "", "suggested_name": it.path.stem,
                             "sensitive": False, "confidence": 10}
                    if not it.date:
                        it.date, it.date_src = dt.date.fromtimestamp(it.mtime).isoformat(), "дата изменения"
                    return it, False

            processed = [it for it in items if it.ai]
            with ThreadPoolExecutor(WORKERS) as ex:
                futs = [ex.submit(work, it) for it in todo]
                for fu in as_completed(futs):
                    it, ok = fu.result()
                    if ok is None:
                        continue
                    processed.append(it)
                    finish(it, ok)
            self.cache.save()
            stopped_early = self.stop.is_set()
            use = [it for it in processed if it.ai]
            emit({"event": "stage", "stage": 2})
            self.ctx = {"scan": sc, "items": use, "root": root, "dest": dest, "model": model, "naming": naming, "wishes": wishes,
                        "partial": stopped_early and len(use) < total, "rules": req.get("rules") or [],
                        "projects_user": req.get("projects") or [], "learned": self.learned}
            clusters = cluster_items(use)
            self.ctx["clusters"] = clusters
            emit({"event": "stage", "stage": 3})
            mapping, projects = plan_mapping(model, wishes, clusters, naming, [p.get("name", "") for p in self.ctx["projects_user"]])
            self.ctx["mapping"] = mapping
            self.ctx["projects"] = projects
            emit({"event": "plan", "plan": build_plan(self.ctx, mapping, naming)})
        except Exception as e:
            traceback.print_exc(file=sys.stderr)
            emit({"event": "error", "message": str(e)})

    # ─── дотюнивание ───
    def tune(self, req: dict) -> dict:
        if not self.ctx:
            raise RuntimeError("План устарел — запустите анализ папки заново")
        ctx = self.ctx
        model = MODELS.get(req.get("model", "accurate"), ctx["model"])
        mapping = dict(req.get("mapping") or ctx["mapping"])
        naming = dict(req.get("naming") or ctx["naming"])
        if "rules" in req:
            ctx["rules"] = req.get("rules") or []
        if "projects" in req:
            ctx["projects_user"] = req.get("projects") or []
        quick = quick_tune(req.get("text", ""), mapping, naming, ctx["clusters"])
        if quick:
            reply, mapping, naming = quick
            plan = build_plan(ctx, mapping, naming)
            ctx["mapping"], ctx["naming"] = mapping, naming
            return {"reply": reply, "question": "", "options": [], "files": plan["files"], "mapping": mapping, "naming": naming}
        hist = req.get("history") or []
        convo = "\n".join(("Пользователь: " if m.get("user") else "Ассистент: ") + m.get("text", "") for m in hist[-8:])
        lines = []
        for cid, g in ctx["clusters"].items():
            years = sorted({i.date[:4] for i in g if i.date and i.date_src != "дата изменения"})
            tops = [t for t, _ in Counter((i.ai.get("topic") or "").lower() for i in g).most_common(2) if t]
            kinds = [k for k, _ in Counter(type_of(i)[0] for i in g).most_common(2)]
            lines.append(f"- id={cid}; путь=«{mapping.get(cid, '')}»; файлов={len(g)}; типы={kinds}; темы={tops}; годы={','.join(years) or 'без даты'}")
        prompt = TUNE_PROMPT.format(wishes=ctx["wishes"] or "(не указаны)", naming=json.dumps({"date_prefix": naming.get("datePrefix", True), "max_words": naming.get("maxWords", 0)}, ensure_ascii=False),
                                    groups="\n".join(lines), convo=convo or "(пусто)", text=req.get("text", ""))
        out, _ = ollama_chat(model, prompt, TUNE_SCHEMA)
        q = (out.get("question") or "").strip()
        if q:
            return {"reply": "", "question": q, "options": [o for o in (out.get("options") or []) if o][:5], "files": None, "mapping": None, "naming": None}
        for mv in out.get("moves") or []:
            cid, path = str(mv.get("id", "")), str(mv.get("path", "")).strip().strip("/")
            if cid in ctx["clusters"] and path:
                mapping[cid] = path
        if isinstance(out.get("date_prefix"), bool):
            naming["datePrefix"] = out["date_prefix"]
        mw = out.get("max_words")
        if isinstance(mw, int) and mw >= 0:
            naming["maxWords"] = mw
        plan = build_plan(ctx, mapping, naming)
        ctx["mapping"], ctx["naming"] = mapping, naming
        return {"reply": out.get("reply") or "Понял, поправил план.", "question": "", "options": [], "files": plan["files"], "mapping": mapping, "naming": naming}

    # ─── проверка на одной картинке ───
    def describe_image(self, req: dict) -> dict:
        model = MODELS.get(req.get("model", "accurate"), MODELS["accurate"])
        if req.get("b64"):
            b64 = req["b64"]
        else:
            img = Image.open(req["path"])
            img = ImageOps.exif_transpose(img)
            b64 = to_b64(img, 896, "JPEG", 85)
        prompt = ("Посмотри на изображение. Опиши его по-русски одним-двумя предложениями, конкретно: что за приложение или документ, "
                  "суммы, получатель, даты — если видны. Затем предложи понятное имя файла (2-7 слов, без расширения; для чеков и оплат — с месяцем и годом в конце, "
                  "например «Оплата Термоленд — авг 2026»). Ответь JSON.")
        schema = {"type": "object", "properties": {"description": {"type": "string"}, "suggested_name": {"type": "string"}}, "required": ["description", "suggested_name"]}
        out, raw = ollama_chat(model, prompt, schema, images=[b64])
        total = raw.get("total_duration", 0) / 1e9
        load = raw.get("load_duration", 0) / 1e9
        return {"desc": out.get("description", "").strip(), "name": clean_name(out.get("suggested_name", "Картинка")), "seconds": round(max(0.3, total - load), 2)}


TYPE_DIRS = {"документы", "скриншоты", "макеты", "дизайн", "картинки", "изображения", "фото", "прототипы", "docs", "documents",
             "screenshots", "images", "design", "mockups", "файлы", "материалы"}

YEAR_PREFIX = re.compile(r"^(\d{4}|\{year\}|\{\{year\}\})\s*[—–-]\s*(.+)$")


def strip_years(path: str) -> str:
    """«Фото/2026 — Вечеринка» → «Фото/Вечеринка», «Фото/{year}» → «Фото»."""
    out = []
    for s in path.split("/"):
        if s in ("{year}", "{{year}}") or re.fullmatch(r"\d{4}", s):
            continue
        m = YEAR_PREFIX.match(s)
        out.append(m.group(2).strip() if m else s)
    return "/".join(out) or path


def quick_tune(text: str, mapping: dict, naming: dict, clusters: dict) -> tuple[str, dict, dict] | None:
    """Частые просьбы (чипы панели «Дотюнить») выполняем сами, без модели — за секунды."""
    t = text.lower()
    mapping, naming = dict(mapping), dict(naming)
    if any(w in t for w in ("корот", "короч", "кратк", "short")) and ("им" in t or "назв" in t or "name" in t):
        cur = int(naming.get("maxWords") or 0)
        n = 4 if not cur else max(2, cur - 1)
        naming["maxWords"] = n
        return f"Сократил имена до {n} {plural(n, ('слова', 'слов', 'слов'))}. Дату в начале имени у фото оставил.", mapping, naming
    if "год" in t and any(w in t for w in ("убер", "убр", "без ", "не нуж")) and ("уров" in t or "папк" in t):
        changed = {cid: strip_years(p) for cid, p in mapping.items() if strip_years(p) != p}
        if changed:
            mapping.update(changed)
            return "Убрал уровень с годами: фото лежат сразу по событиям.", mapping, naming
    if "без даты" in t and "фото" in t and any(w in t for w in ("одну", "вместе", "сложи", "собери")):
        for cid, g in clusters.items():
            if g and g[0].category == "photo_nodate":
                mapping[cid] = "Фото/Без даты"
        return "Фото без даты сложил в одну папку «Без даты».", mapping, naming
    if "дат" in t and "им" in t and any(w in t for w in ("убер", "убр", "без", "не нуж", "не став")):
        naming["datePrefix"] = False
        return "Убрал дату из начала имён фото.", mapping, naming
    return None


RULE_DIR = {"installers": "Установщики", "archives": "Архивы", "temp": "Временные файлы", "service": "Служебные файлы", "other": "Разобрать вручную"}
RULE_DESC = {"installers": "Установщик программы.", "archives": "Архив.", "temp": "Недокачанный или временный файл.",
             "service": "Служебный файл.", "other": "Файл неизвестного типа."}


def type_of(it: Item) -> tuple[str, str, str]:
    """(тип для фильтров, иконка, цвет)"""
    k = it.ai.get("kind", "")
    if it.kind == "rule":
        return "другое", {"installers": "install_desktop", "archives": "folder_zip", "temp": "hourglass_empty"}.get(it.rule, "draft"), "#8b8b90"
    if it.kind == "video":
        return "видео", "movie", COL["b"]
    if it.kind == "audio":
        return "аудио", "music_note", COL["b"]
    if it.kind == "image":
        if it.ai.get("sensitive") or k == "скан документа":
            return "скан", "badge", COL["t"]
        if k == "фото":
            return "фото", "photo_camera", COL["p"]
        if k == "скриншот":
            if "чек" in (it.ai.get("topic", "") + it.ai.get("description", "")).lower() or "оплат" in it.ai.get("description", "").lower():
                return "скриншот", "receipt_long", COL["t"]
            return "скриншот", "screenshot_monitor", COL["v"]
        return "картинка", "image", COL["p"]
    if k == "счёт/чек/финансы":
        return "документ", "receipt_long", COL["t"]
    if it.ai.get("sensitive") or k == "личный документ":
        return "скан" if it.ext == ".pdf" and it.meta.get("pages") else "документ", "badge", COL["t"]
    return "документ", "description", COL["t"]


def thumb_label(it: Item) -> str:
    t = type_of(it)[0]
    if it.ext == ".pdf":
        p = it.meta.get("pages")
        return f"pdf, {p} стр." if p else "pdf"
    if it.kind == "doc":
        return it.ext.lstrip(".")
    return t


def themes_of(topics: Counter, tk: dict) -> list[dict]:
    out = []
    for key, n in topics.most_common(8):
        if n < 2:
            continue
        kinds = tk[key]
        label = next((k[9:] for k in kinds if k.startswith("__label__")), key)
        real = Counter({k: v for k, v in kinds.items() if not k.startswith("__label__")})
        major = real.most_common(1)[0][0] if real else "документ"
        c = {"фото": COL["p"], "скриншот": COL["v"], "картинка": COL["p"], "видео": COL["b"], "аудио": COL["b"]}.get(major, COL["t"])
        out.append({"label": label, "n": f"{n}" + (" фото" if major == "фото" and len(real) == 1 else ""), "c": c})
    return out[:7]


BAD = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def clean_name(name: str, limit: int = 90) -> str:
    name = BAD.sub(" ", str(name)).replace("_", " ")
    name = re.sub(r"\s+", " ", name).strip(" .")
    return name[:limit].rstrip(" .") or "файл"


def ui_file(it: Item, naming: dict) -> dict:
    t, icon, c = type_of(it)
    return {"id": it.id, "old": it.path.name, "thumb": it.thumb, "thumbLabel": thumb_label(it), "kind": t.capitalize(),
            "desc": it.ai.get("description", ""), "topic": it.ai.get("topic", ""), "name": proposed_name(it, naming), "icon": icon, "c": c}


TAIL_WORDS = {"в", "во", "на", "за", "с", "со", "и", "а", "но", "по", "для", "к", "ко", "у", "о", "об", "от", "из", "под", "над", "при",
              "про", "без", "до", "через", "или", "the", "a", "an", "of", "for", "in", "on", "at", "to", "and", "with", "—", "-", "–"}
ADJ_END = ("ой", "ый", "ий", "ая", "яя", "ое", "ее", "ые", "ие", "ых", "их", "ым", "им", "ую", "юю", "ого", "его")
DATE_TAIL = re.compile(r"\s*[—–-]\s*(янв|фев|мар|апр|ма[йя]|июн|июл|авг|сен|окт|ноя|дек)[а-я]*\.?\s+\d{4}$", re.I)


def shorten(stem: str, n: int) -> str:
    """Обрезает имя до n слов, не оставляя на конце предлог или повисшее прилагательное."""
    m = DATE_TAIL.search(stem)
    tail = m.group(0) if m else ""
    body = stem[: m.start()] if m else stem
    words = body.split()
    if len(words) <= n:
        return stem
    words = words[:n]
    while len(words) > 1 and (words[-1].lower().strip(",.") in TAIL_WORDS or (len(words) > 2 and words[-1].lower().endswith(ADJ_END))):
        words.pop()
    return (" ".join(words).rstrip(" —-,") + tail) or stem


def proposed_name(it: Item, naming: dict, base: str | None = None) -> str:
    if not naming.get("rename", True) or it.kind == "rule":
        return it.path.name
    stem = clean_name(base or it.ai.get("suggested_name") or it.path.stem)
    mw = int(naming.get("maxWords") or 0)
    if mw > 0:
        stem = shorten(stem, mw)
    is_photo = it.kind == "image" and it.ai.get("kind") == "фото"
    if is_photo and naming.get("datePrefix", True) and it.date_src in ("EXIF", "имя файла") and not stem.startswith(it.date):
        stem = f"{it.date} {stem}"
    return stem + it.ext


# ───────────────────────────── группировка ─────────────────────────────

def category_of(it: Item) -> str:
    if it.kind == "rule":
        return "rule:" + it.rule
    if it.kind == "image":
        k = it.ai.get("kind")
        if it.ai.get("sensitive") or k == "скан документа":
            return "doc"
        if k == "фото" and (it.meta.get("camera") or it.date_src == "EXIF"):
            return "photo" if it.date_src == "EXIF" else "photo_nodate"
        return "image"
    return it.kind


def time_events(items: list[Item], gap_days: float = 2.0) -> list[list[Item]]:
    items = sorted(items, key=lambda x: x.date)
    groups, cur, prev = [], [], None
    for it in items:
        d = dt.date.fromisoformat(it.date)
        if prev and (d - prev).days > gap_days:
            groups.append(cur)
            cur = []
        cur.append(it)
        prev = d
    if cur:
        groups.append(cur)
    return groups


def agglomerative(vecs: np.ndarray, threshold: float) -> list[int]:
    """Иерархическая кластеризация, average linkage по косинусному расстоянию (векторы нормированы).
    Сливает две ближайшие группы, пока расстояние между ними меньше порога. Без sklearn/scipy — движок легче."""
    n = len(vecs)
    if n <= 1:
        return [0] * n
    d = 1.0 - (vecs @ vecs.T).astype(np.float64)
    np.fill_diagonal(d, np.inf)
    size = np.ones(n)
    members = [[i] for i in range(n)]
    while True:
        k = int(np.argmin(d))
        i, j = divmod(k, n)
        if not np.isfinite(d[i, j]) or d[i, j] >= threshold:
            break
        merged = (size[i] * d[i] + size[j] * d[j]) / (size[i] + size[j])
        d[i, :] = merged
        d[:, i] = merged
        d[i, i] = np.inf
        d[j, :] = np.inf
        d[:, j] = np.inf
        size[i] += size[j]
        members[i] += members[j]
        members[j] = []
    labels = [0] * n
    for cid, mem in enumerate(m for m in members if m):
        for p in mem:
            labels[p] = cid
    return labels


def semantic(items: list[Item], vecs: np.ndarray, threshold: float = 0.32) -> list[list[Item]]:
    if len(items) <= 1:
        return [items]
    labels = agglomerative(vecs, threshold)
    g = defaultdict(list)
    for it, lb in zip(items, labels):
        g[lb].append(it)
    return list(g.values())


def embed_text(i: Item) -> str:
    return f"{i.ai.get('kind', '')}. Тема: {i.ai.get('topic', '')}. {i.ai.get('description', '')} {i.ai.get('text_on_image', '')} Файл: {i.path.stem}"[:700]


def cluster_items(items: list[Item]) -> dict[str, list[Item]]:
    by_cat = defaultdict(list)
    for it in items:
        it.category = category_of(it)
        by_cat[it.category].append(it)
    need = [it for it in items if it.category not in ("photo",) and not it.category.startswith("rule:")]
    vecs: dict[int, np.ndarray] = {}
    if need:
        texts = [embed_text(i) for i in need]
        try:
            arr = embed(texts)
            vecs = {it.id: v for it, v in zip(need, arr)}
            for it, v in zip(need, arr):
                it.vec = v
        except Exception:
            vecs = {}
    clusters: dict[str, list[Item]] = {}
    n = 0
    for cat, its in by_cat.items():
        if cat == "photo":
            groups = time_events(its)
        elif cat.startswith("rule:"):
            groups = [its]
        elif vecs and all(i.id in vecs for i in its):
            # в небольших папках группируем почти пофайлово — план точнее; в больших — крупнее, чтобы модель справилась
            thr = 0.2 if len(items) <= 80 else 0.3
            groups = semantic(its, np.stack([vecs[i.id] for i in its]), thr)
        else:
            tg = defaultdict(list)
            for i in its:
                tg[(i.ai.get("topic") or "").lower()].append(i)
            groups = list(tg.values())
        for g in groups:
            n += 1
            cid = f"c{n}"
            for it in g:
                it.cluster = cid
            clusters[cid] = g
    return clusters


VERSION_RX = re.compile(r"^(?P<base>.+?)(?:\s*\((?P<n>\d{1,3})\)|[_\- ](?:v|копия|copy)?(?P<m>\d{1,2})|[_\- ](?:копия|copy))$", re.I)
CAMERA_RX = re.compile(r"^(img|dsc|vid|pxl|image|photo|screenshot|снимок|scan|scan\d|unnamed|document|документ|новый документ|file)[\s_\-]?", re.I)


def version_groups(items: list[Item]) -> dict[int, tuple[list[Item], Item]]:
    """Находит несколько версий одного файла: 'макет (2).png', 'макет_3.png' ... → {id: (группа, последняя)}."""
    fam = defaultdict(list)
    for it in items:
        stem = it.path.stem.strip()
        base = stem
        for _ in range(2):
            m = VERSION_RX.match(base)
            if not m:
                break
            base = m.group("base").strip()
        if CAMERA_RX.match(base) or len(base) < 5:
            continue
        fam[(base.lower(), it.ext)].append(it)
    out = {}
    for (_, _), g in fam.items():
        if len(g) < 2:
            continue
        latest = max(g, key=lambda x: x.mtime)
        for it in g:
            out[it.id] = (g, latest)
    return out


# ───────────────────────────── план ─────────────────────────────

PLAN_SCHEMA = {
    "type": "object",
    "properties": {
        "projects": {"type": "array", "items": {"type": "string"}},
        "assignments": {"type": "array", "items": {"type": "object", "properties": {"id": {"type": "string"}, "path": {"type": "string"}}, "required": ["id", "path"]}},
    },
    "required": ["projects", "assignments"],
}

PLAN_PROMPT = """Ты проектируешь структуру папок для личных файлов пользователя.
Пожелания пользователя: «{wishes}»
{langnote}
Ниже группы похожих файлов: id, категория, количество, годы, темы, типы, примеры описаний и исходных имён.
Для КАЖДОГО id верни путь целевой папки относительно корня, через «/».
Правила:
- Следуй пожеланиям пользователя. Глубина 2-3 уровня, названия папок короткие и человеческие. Лучше меньше папок, чем больше.
- Общие корневые разделы: «Проекты», «Фото», «Документы», «Скриншоты», «Видео», «Музыка». Не плоди синонимы.
- Проект — это то, над чем пользователь САМ работает: его приложение, сайт, бренд (макеты, прототипы, ТЗ, спецификации, промпты, скриншоты интерфейсов в разработке). Для каждого такого проекта сделай «Проекты/<Название>» и перечисли их в поле projects.
- Разные проекты — разные папки. Группа относится к проекту, только если его название явно есть в темах, описаниях или именах файлов группы.
- НЕ проекты: игры, фильмы, сторонние приложения и сайты, которыми пользователь просто пользуется. Скриншоты игр — «Скриншоты/Игры», прочие скриншоты — «Скриншоты/<тема>».
- Внутри проекта не делай подпапок по типу файлов («Документы», «Скриншоты») — клади файлы прямо в «Проекты/<Название>».
- Документы (pdf, docx, md, txt, html) кладутся только в «Документы/…» или «Проекты/…» — не в «Фото» и не в «Скриншоты».
- Фото с камеры — по годам и событиям («Фото/2026 — Вечеринка у Ани»). Для группы фото за разные годы используй плейсхолдер {{year}} — он подставится по дате каждого файла.
- Личные документы (паспорт, СНИЛС, удостоверения) — «Документы/Личные документы».
- Чеки, счета, оплаты, договоры — «Документы/Финансы». Учёба — «Документы/Учёба».
- Если группа непонятная или смешанная — «Разобрать вручную».
- Группы про одно и то же получают один и тот же путь.

{groups}"""

TUNE_SCHEMA = {
    "type": "object",
    "properties": {
        "reply": {"type": "string"},
        "question": {"type": "string"},
        "options": {"type": "array", "items": {"type": "string"}},
        "moves": {"type": "array", "items": {"type": "object", "properties": {"id": {"type": "string"}, "path": {"type": "string"}}, "required": ["id", "path"]}},
        "date_prefix": {"type": "boolean"},
        "max_words": {"type": "integer"},
    },
    "required": ["reply", "question", "options", "moves", "date_prefix", "max_words"],
}

TUNE_PROMPT = """Ты помогаешь поправить план сортировки файлов. Файлы уже проанализированы — меняется только план.
Исходные пожелания: «{wishes}»
Текущие настройки имён: {naming} (date_prefix — дата в начале имени у фото; max_words — сколько слов оставлять в имени, 0 — без ограничения).
Группы файлов и их текущие папки:
{groups}

Разговор:
{convo}

Новая просьба пользователя: «{text}»

Ответь JSON:
- Если просьба неоднозначна и без уточнения можно сделать не то (например, непонятно, какие именно папки считать проектами) — задай короткий вопрос в question и дай 2-4 варианта ответа в options; moves оставь пустым.
- Иначе: question — пустая строка; moves — только группы, у которых меняется папка (id и новый путь, плейсхолдер {{year}} допустим); date_prefix и max_words — новые настройки имён (если просьба их не касается — верни текущие);
  reply — одно-два предложения по-русски: что ты поменял, в стиле «Сложил скриншоты в папки проектов. Остальное не трогал.»"""


BRAND_RX = re.compile(r"\b([A-Z][A-Za-z0-9]{2,}(?:[ :\-][A-Z][A-Za-z0-9]{1,})?)\b")
QUOTED_RX = re.compile(r"«([^»]{3,30})»")
BRAND_STOP = {"pdf", "html", "spa", "png", "jpg", "jpeg", "img", "dsc", "vid", "ui", "ux", "api", "id", "it", "ai", "ok", "usb", "gui", "pc",
              "cpu", "gpu", "ru", "en", "telegram", "windows", "android", "iphone", "apple", "google", "microsoft", "excel", "word",
              "powerpoint", "steam", "discord", "nvidia", "bundled", "page", "heic", "mp3", "mp4", "esim", "sim", "the", "and", "for",
              "new", "app", "web", "ios", "wifi", "vpn", "url", "email", "sms", "qr", "faq", "css", "json", "docx", "screenshot",
              "image", "photo", "document", "file", "unnamed", "react", "figma", "chrome", "edge", "youtube", "whatsapp", "instagram"}


def norm_key(s: str) -> str:
    return re.sub(r"[\s:\-_.,]+", "", s).lower()


def it_text(it: Item) -> str:
    return norm_key(" ".join([it.ai.get("topic", ""), it.ai.get("suggested_name", ""), it.ai.get("description", ""),
                              it.ai.get("text_on_image", ""), it.path.stem]))


def own_brands(it: Item) -> list[str]:
    """Собственные названия файла (бренд, приложение, компания) — по имени, теме и предложенному имени."""
    text = " ".join([it.ai.get("suggested_name", ""), it.ai.get("topic", ""), it.path.stem.replace("_", " ")])
    out = []
    for c in BRAND_RX.findall(text):
        c = strip_stop(c)
        if c and norm_key(c) not in BRAND_STOP and len(norm_key(c)) >= 4 and not norm_key(c).isdigit() and c not in out:
            out.append(c)
    return out


def strip_stop(c: str) -> str:
    """«SPA Termoland» → «Termoland»: служебные слова внутри составного названия отбрасываем."""
    parts = [p for p in re.split(r"[\s:\-]+", c.strip()) if p and p.lower() not in BRAND_STOP]
    return " ".join(parts)


def project_hints(items: list[Item]) -> list[tuple[str, int]]:
    """Названия, которые повторяются в разных файлах: кандидаты в проекты пользователя (Termoland, Klutz, SIM PIE…)."""
    seen: dict[str, set[int]] = defaultdict(set)
    forms: dict[str, Counter] = defaultdict(Counter)
    gamey: Counter = Counter()
    for it in items:
        if it.kind not in ("image", "doc"):
            continue
        text = " ".join([it.ai.get("topic", ""), it.ai.get("suggested_name", ""), it.path.stem.replace("_", " ").replace("-", " ")])
        cands = BRAND_RX.findall(text) + QUOTED_RX.findall(text + " " + it.ai.get("description", ""))
        is_game = "игр" in (it.ai.get("description", "") + it.ai.get("topic", "")).lower()
        for c in cands:
            c = strip_stop(c)
            key = re.sub(r"[\s:\-]+", "", c).lower()
            if key in BRAND_STOP or len(key) < 4 or key.isdigit():
                continue
            seen[key].add(it.id)
            forms[key][c.strip()] += 1
            if is_game:
                gamey[key] += 1
    out = []
    for key, ids in seen.items():
        if len(ids) >= 2 and gamey[key] * 2 < len(ids):
            out.append((forms[key].most_common(1)[0][0], len(ids)))
    out.sort(key=lambda x: -x[1])
    return out[:8]


def plan_mapping(model: str, wishes: str, clusters: dict[str, list[Item]], naming: dict,
                 user_projects: list[str] | None = None) -> tuple[dict[str, str], list[str]]:
    all_items = [i for g in clusters.values() for i in g]
    mine = [p.strip() for p in user_projects or [] if p and p.strip()]
    hints = [(n, k) for n, k in project_hints(all_items) if norm_key(n) not in {norm_key(p) for p in mine}]
    hint_line = ("Подсказка: эти названия повторяются в нескольких файлах — возможно, это проекты пользователя (используй их как названия папок проектов, если это действительно проекты): "
                 + ", ".join(f"{n} ({k} файла)" for n, k in hints)) if hints else ""
    if mine:
        root = "Projects" if naming.get("lang") == "en" else "Проекты"
        hint_line = (f"Проекты пользователя (он сам их назвал): {', '.join(mine)}. Файлы этих проектов клади в «{root}/<название>», "
                     f"название пиши ровно так, как здесь.\n" + hint_line).strip()
    lines = []
    for cid, g in clusters.items():
        cat = g[0].category
        if cat.startswith("rule:"):
            continue
        years = sorted({i.date[:4] for i in g if i.date_src in ("EXIF", "имя файла", "текст документа", "текст на картинке")})
        tops = [t for t, _ in Counter((i.ai.get("topic") or "").lower() for i in g).most_common(5) if t]
        kinds = [k for k, _ in Counter(i.ai.get("kind", "") for i in g).most_common(2)]
        samples = [i.ai.get("description", "")[:110] for i in g[:3]]
        names = [i.path.name[:40] for i in g[:3]]
        what = "документы" if cat == "doc" else ("фото с камеры" if cat.startswith("photo") else ("картинки и скриншоты" if cat == "image" else cat))
        lines.append(f"- id={cid}; что={what}; файлов={len(g)}; годы={','.join(years) or 'без даты'}; типы={kinds}; темы={tops}; примеры={samples}; имена={names}")
    mapping: dict[str, str] = {}
    projects: list[str] = []
    langnote = "Названия папок — по-английски." if naming.get("lang") == "en" else "Названия папок — по-русски."
    chunk = 30
    try:
        for i in range(0, len(lines), chunk):
            part = "\n".join(lines[i:i + chunk])
            if mapping:
                part += "\n\nУже созданные папки (переиспользуй, если подходит): " + "; ".join(sorted(set(mapping.values()))[:80])
            if hint_line:
                part = hint_line + "\n\n" + part
            out, _ = ollama_chat(model, PLAN_PROMPT.format(wishes=wishes or "(не указаны — предложи разумную структуру)", langnote=langnote, groups=part), PLAN_SCHEMA)
            for a in out.get("assignments", []):
                p = str(a.get("path", "")).strip().strip("/")
                if p:
                    mapping[str(a.get("id"))] = p
            projects += [p.strip() for p in out.get("projects", []) if p and p.strip()]
    except Exception:
        traceback.print_exc(file=sys.stderr)
    # запасной вариант для групп без ответа модели
    for cid, g in clusters.items():
        if cid in mapping:
            continue
        cat = g[0].category
        if cat.startswith("rule:"):
            mapping[cid] = RULE_DIR[cat[5:]]
        elif cat == "photo":
            mapping[cid] = "Фото/{year}"
        elif cat == "photo_nodate":
            mapping[cid] = "Фото/Без даты"
        elif cat == "image":
            mapping[cid] = "Скриншоты" if any(i.ai.get("kind") == "скриншот" for i in g) else "Картинки"
        elif cat == "video":
            mapping[cid] = "Видео"
        elif cat == "audio":
            mapping[cid] = "Музыка"
        else:
            mapping[cid] = "Документы"
    # проекты — только те, для которых реально есть папки
    tops = {p.split("/")[1] for p in mapping.values() if p.startswith(("Проекты/", "Projects/")) and len(p.split("/")) > 1}
    seen, proj = set(), []
    for p in projects + sorted(tops):
        if p in tops and p.lower() not in seen:
            seen.add(p.lower()); proj.append(p)
    return mapping, proj[:6]


YEAR_DUP = re.compile(r"^(\d{4})\s*[—–-]\s*\1$")
YEAR_NODATE = re.compile(r"^\d{4}\s*[—–-]\s*(без даты|no date)$", re.I)
YEAR_SEG = re.compile(r"^\d{4}(\s*[—–-].*)?$")
MANUAL = {"разобрать вручную", "sort manually"}


def is_unreliable_photo(it: Item) -> bool:
    return it.kind == "image" and it.ai.get("kind") == "фото" and it.date_src == "дата изменения"


def tidy_path(path: str, it: Item, en: bool) -> str:
    """Чистит путь, который предложила модель: дубли года, «Без даты» с годом, вложенный «Разобрать вручную»."""
    segs = []
    for s in (x.strip() for x in path.split("/")):
        if not s:
            continue
        m = YEAR_DUP.match(s)
        if m:
            s = m.group(1)
        if YEAR_NODATE.match(s):
            s = "No date" if en else "Без даты"
        segs.append(clean_name(s, 60))
    if not segs or any(s.lower() in MANUAL for s in segs):
        return "Sort manually" if en else "Разобрать вручную"
    if is_unreliable_photo(it) and segs[0] in ("Фото", "Photos"):
        # без надёжной даты фото не раскладываем по годам и событиям
        return f"{segs[0]}/" + ("No date" if en else "Без даты")
    return "/".join(segs)


# ───────────────────── правила, проекты и выученное: решают раньше модели ─────────────────────

LEARN_SIM = 0.80  # насколько файл должен быть похож на запомненный (bge-m3: чужие ~0.5–0.72, одно событие ~0.8+, копии ~0.97)
WORD = "0-9a-zа-я"
TYPE_RULE = {"photo": "Фото", "screenshot": "Скриншоты", "image": "Картинки", "doc": "Документы", "pdf": "PDF",
             "table": "Таблицы", "video": "Видео", "audio": "Аудио"}


def norm_text(s: str) -> str:
    s = re.sub(r"([a-zа-я])([A-ZА-Я])", r"\1 \2", s)  # SimPie → sim pie
    s = s.lower().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-я*?]+", " ", s).strip()


def kw_hit(kw: str, text: str) -> bool:
    """Ключевое слово в тексте: с начала слова, окончания любые («термоленд» найдёт «Термоленда»);
    короткие слова — только целиком; длинные — и без пробелов («sim pie» = «SimPie»)."""
    k = norm_text(kw).replace("*", " ").replace("?", " ").strip()
    if len(k) < 2:
        return False
    flat = k.replace(" ", "")
    if len(flat) >= 5 and flat in text.replace(" ", ""):
        return True
    tail = "" if len(k) >= 4 else f"(?![{WORD}])"
    return re.search(f"(?<![{WORD}])" + re.escape(k) + tail, text) is not None


def about_text(it: Item) -> str:
    """О чём файл: имя, папка, тема и описание от модели, текст на картинке. По нему узнаём проекты:
    случайное упоминание проекта где-то в теле документа сюда не попадает."""
    return norm_text(" ".join([it.path.stem, it.from_rel, it.ai.get("topic", ""), it.ai.get("suggested_name", ""),
                               it.ai.get("description", ""), it.ai.get("text_on_image", "")]))


def content_text(it: Item) -> str:
    """Всё, что известно о содержимом, включая текст документа, — для правил «Текст содержит»."""
    if it.kind == "doc" and "text" not in it.meta:
        try:  # описание пришло из кэша без текста — дочитываем
            it.meta["text"] = doc_text(it, 1500)
        except Exception:
            it.meta["text"] = ""
    return about_text(it) + " " + norm_text(it.meta.get("text", ""))


def type_hit(value: str, it: Item) -> bool:
    v = value.strip().lower()
    tp = type_of(it)[0]
    if v == "photo":
        return tp == "фото"
    if v == "screenshot":
        return tp == "скриншот"
    if v == "image":
        return it.kind == "image"
    if v == "doc":
        return it.kind == "doc" or tp == "скан"
    if v == "pdf":
        return it.ext == ".pdf"
    if v == "table":
        return it.ext in (".xlsx", ".xls", ".csv", ".ods")
    if v in ("video", "audio"):
        return it.kind == v
    exts = {"." + e.strip(" .*") for e in re.split(r"[,\s;]+", v) if e.strip(" .*")}
    return it.ext in exts


def rule_hit(rule: dict, it: Item, cache: dict) -> bool:
    field_, value = rule.get("field", "name"), str(rule.get("value", "")).strip()
    if not value:
        return False
    if field_ == "type":
        return type_hit(value, it)
    if field_ == "name":
        if "*" in value or "?" in value:
            import fnmatch
            return any(fnmatch.fnmatch(it.path.name.lower(), v.strip().lower()) for v in value.split(",") if v.strip())
        text = norm_text(it.path.name)
    elif field_ == "from":
        text = norm_text(it.from_rel + " " + str(it.path.parent))
    else:
        if it.id not in cache:
            cache[it.id] = content_text(it)
        text = cache[it.id]
    return any(kw_hit(v, text) for v in value.split(",") if v.strip())


def overrides(ctx: dict, items: list[Item], en: bool) -> dict[int, dict]:
    """Куда файл кладётся без модели: правила пользователя → его проекты → выученные примеры."""
    rules = [r for r in ctx.get("rules") or [] if str(r.get("folder", "")).strip() and str(r.get("value", "")).strip()]
    projects = [p for p in ctx.get("projects_user") or [] if str(p.get("name", "")).strip()]
    learned = ctx.get("learned") or []
    out: dict[int, dict] = {}
    cache: dict[int, str] = {}
    lv = np.stack([e["v"] for e in learned]) if learned else None
    proj_root = "Projects" if en else "Проекты"
    for it in items:
        if it.kind == "rule":
            continue
        for r in rules:
            if rule_hit(r, it, cache):
                what = {"name": "имя", "text": "текст", "from": "откуда", "type": "тип"}.get(r.get("field"), "")
                val = TYPE_RULE.get(r["value"], r["value"]) if r.get("field") == "type" else r["value"]
                out[it.id] = {"path": r["folder"], "rename": r.get("rename") or "auto", "src": "rule",
                              "why": f"Ваше правило: {what} — «{val}» → «{r['folder'].strip().strip('/').replace('{year}', 'год')}»."}
                break
        if it.id in out:
            continue
        if projects:
            text = about_text(it)
            for p in projects:
                words = [p["name"]] + [k for k in p.get("keywords") or [] if str(k).strip()]
                hit = next((w for w in words if kw_hit(str(w), text)), None)
                if hit:
                    out[it.id] = {"path": f"{proj_root}/{p['name'].strip()}", "rename": "auto", "src": "project",
                                  "why": f"Ваш проект «{p['name'].strip()}»: в файле есть «{str(hit).strip()}»."}
                    break
        if it.id in out or lv is None or it.vec is None:
            continue
        sims = lv @ np.asarray(it.vec, dtype=np.float32)
        order = np.argsort(-sims)
        k = int(order[0])
        rival = next((float(sims[j]) for j in order[1:] if learned[j]["folder"] != learned[k]["folder"]), 0.0)
        if sims[k] >= LEARN_SIM and sims[k] - rival > 0.02:  # два примера из разных папок почти одинаково близки — не угадываем
            e = learned[k]
            why = (f"Вы переложили этот файл в «{e['folder']}» и попросили запомнить." if e["label"] == it.path.name
                   else f"Похож на «{e['label']}» — его вы переложили в «{e['folder']}».")
            out[it.id] = {"path": e["folder"], "rename": "auto", "src": "learned", "why": why}
    for o in out.values():
        o["path"] = "/".join(clean_name(s, 60) for s in str(o["path"]).replace("\\", "/").split("/") if s.strip() and s.strip() not in (".", ".."))
    return {k: v for k, v in out.items() if v["path"]}


def build_plan(ctx: dict, mapping: dict[str, str], naming: dict) -> dict:
    items: list[Item] = ctx["items"]
    clusters: dict[str, list[Item]] = ctx["clusters"]
    sc: Scan = ctx["scan"]
    root: Path = ctx["root"]
    dest: Path = ctx["dest"]
    en = naming.get("lang") == "en"
    versions = version_groups(items)
    forced = overrides(ctx, items, en)
    files = []
    taken: set[str] = set()

    def unique(folder: str, name: str) -> str:
        stem, ext = os.path.splitext(name)
        cand, k = name, 2
        while f"{folder}/{cand}".lower() in taken:
            cand = f"{stem} ({k}){ext}"
            k += 1
        taken.add(f"{folder}/{cand}".lower())
        return cand

    # сколько файлов в каждом проекте — маленьким проектам не нужны подпапки по типу
    proj_size: Counter = Counter()
    for it in items:
        p = mapping.get(it.cluster) or ""
        if p.startswith(("Проекты/", "Projects/")) and len(p.split("/")) > 1:
            proj_size[clean_name(p.split("/")[1], 60)] += 1

    # проход 1: куда кладём каждый файл
    target: dict[int, str] = {}
    for it in items:
        year = it.date[:4] if it.date and not is_unreliable_photo(it) else ""
        nodate = "No date" if en else "Без даты"
        if it.id in forced:  # правило, проект или выученный пример — как сказал пользователь
            target[it.id] = forced[it.id]["path"].replace("{{year}}", year or nodate).replace("{year}", year or nodate)
            continue
        path = mapping.get(it.cluster) or "Разобрать вручную"
        path = path.replace("{{year}}", year or nodate).replace("{year}", year or nodate)
        path = tidy_path(path, it, en)
        parts = path.split("/")
        # файл попадает в проект, только если название проекта есть в его имени, теме или описании
        if len(parts) >= 2 and parts[0] in ("Проекты", "Projects"):
            own = it_text(it)
            if norm_key(parts[1]) not in own:
                brands = own_brands(it)
                if brands:
                    parts[1] = brands[0]
        # в небольших проектах — без подпапок по типу файлов
        if len(parts) >= 3 and parts[0] in ("Проекты", "Projects") and proj_size.get(parts[1], 0) < 15 and parts[2].lower() in TYPE_DIRS:
            parts = parts[:2] + parts[3:]
        # документы не кладём в «Фото» и «Скриншоты»
        if it.kind == "doc" and parts[0] in ("Скриншоты", "Фото", "Screenshots", "Photos", "Картинки"):
            parts = ["Documents" if en else "Документы"]
        path = "/".join(parts)
        if it.ai.get("sensitive") and it.ai.get("kind") in ("скан документа", "личный документ") and not path.startswith(("Документы/Личные", "Documents/Personal")):
            path = "Documents/Personal" if en else "Документы/Личные документы"
        conf = int(max(0, min(100, it.ai.get("confidence", 70) or 70)))
        if conf < 40:
            path = "Sort manually" if en else "Разобрать вручную"
        target[it.id] = path

    # проход 2: папки из одного файла сворачиваем в родительскую (кроме проектов, событий с фото и личных документов)
    counts = Counter(target.values())
    for iid, p in list(target.items()):
        segs = p.split("/")
        keep = segs[0] in ("Проекты", "Projects", "Фото", "Photos") or p.startswith(("Документы/Личные", "Documents/Personal")) or iid in forced
        if counts[p] == 1 and len(segs) >= 2 and not keep:
            target[iid] = "/".join(segs[:-1])

    # нумерация фото внутри событий
    event_idx: dict[int, tuple[int, int]] = {}
    for cid, g in clusters.items():
        if g and g[0].category == "photo" and len(g) >= 3:
            for k, it in enumerate(sorted(g, key=lambda x: (x.date, x.path.name)), 1):
                event_idx[it.id] = (k, len(g))

    for it in sorted(items, key=lambda x: x.rel.lower()):
        tp, icon, _c = type_of(it)
        conf = int(max(0, min(100, it.ai.get("confidence", 70) or 70)))
        check = None
        reason = (it.ai.get("description") or "").strip()
        path = target[it.id]
        fo = forced.get(it.id)
        if path.startswith(("Документы/Личные", "Documents/Personal")) and it.ai.get("sensitive"):
            reason += " Личные данные — в отдельную папку, превью размыто."
        name = proposed_name(it, naming)
        if fo and fo["rename"] == "keep":
            name = it.path.name
        elif fo and fo["rename"] == "date":
            stem = os.path.splitext(name)[0]
            name = (stem if stem.startswith(it.date) else f"{it.date} {stem}") + it.ext
        if it.id in event_idx and naming.get("rename", True) and not (fo and fo["rename"] != "auto"):
            k, n = event_idx[it.id]
            leaf = re.sub(r"^\d{4}\s*[—\-–]\s*", "", path.split("/")[-1]).strip() or "Событие"
            if leaf.lower() not in ("без даты", str(it.date[:4])):
                base = f"{leaf} {k:02d}"
                name = (f"{it.date} {base}" if naming.get("datePrefix", True) else base) + it.ext
            reason = f"Одна из {n} фото одного события. " + reason
        if it.id in versions:
            g, latest = versions[it.id]
            if it.id == latest.id:
                check = f"{len(g)} {plural(len(g), ('версия', 'версии', 'версий'))} одного файла — оставили последнюю сверху"
                reason = f"Последняя из {len(g)} версий одного файла. Остальные {len(g) - 1} предлагаю убрать в «Старые версии». " + reason
            else:
                path = path + "/Старые версии"
                reason = f"Старая версия файла «{latest.path.name}» — убираем в «Старые версии»."
        if name_is_broken(it.path.name) and naming.get("rename", True):
            reason += " Испорченное имя — восстановили по содержимому."
        if it.kind in ("video", "audio"):
            gb = it.size / 1024 ** 3
            sz = f"{gb:.1f} ГБ" if gb >= 1 else f"{it.size / 1024 ** 2:.0f} МБ"
            reason = f"{'Видео' if it.kind == 'video' else 'Аудио'} {sz} — разобрано по имени и дате, без анализа содержимого. " + reason
            conf = min(conf, 75)
        if it.kind == "rule":
            reason = f"{RULE_DESC[it.rule]} Вы вернули такие файлы в разбор — кладём в «{RULE_DIR[it.rule]}»."
        date_disp, dsrc = fmt_date(it.date), it.date_src
        if it.meta.get("exif_reset") and it.date_src == "дата изменения":
            d = it.meta["exif_reset"]
            date_disp, dsrc = f"{d[8:10]}.{d[5:7]}.{d[:4]}?", "EXIF, ненадёжно"
            check = check or "Дата ненадёжна: часы камеры были сброшены"
            conf = min(conf, 58)
        if fo:  # решение пользователя — без «модель не уверена»
            reason = fo["why"] + (" " + reason if reason else "")
            conf = max(conf, 97)
        else:
            if conf < 40 and path != "Разобрать вручную":
                path = "Разобрать вручную"
            if path == "Разобрать вручную":
                conf = min(conf, 39)
                check = check or "Модель не уверена, куда положить"
            elif conf < 55:
                check = check or "Модель не уверена, куда положить"
        name = unique(path, name)
        files.append({
            "id": it.id, "abs": str(it.path), "cur": str(it.path), "old": it.path.name, "fromRel": it.from_rel,
            "name": name, "orig": name, "to": path, "engineTo": path, "type": tp, "icon": icon, "reason": reason.strip(),
            "date": date_disp, "dsrc": dsrc, "conf": conf, "check": check, "priv": bool(it.ai.get("sensitive")),
            "thumb": it.thumb, "kind": it.kind if it.kind != "rule" else "doc", "cluster": it.cluster, "size": it.size, "rejected": False,
            "by": fo["src"] if fo else None,
        })
    skipped = sum(sc.skipped_counts.values())
    same = str(dest).rstrip("\\/").lower() == str(root).rstrip("\\/").lower()
    proj_count = Counter(f["to"].split("/")[1] for f in files if f["to"].startswith(("Проекты/", "Projects/")) and len(f["to"].split("/")) > 1)
    projects = [p for p, _ in proj_count.most_common(6)]
    return {
        "root": str(root), "rootDisplay": display_name(root), "dest": str(dest), "destDisplay": display_name(root) if same else display_name(dest),
        "files": files, "locked": [{k: v for k, v in l.items() if k != "cat"} for l in sc.locked], "skipped": skipped,
        "projects": projects, "rootFiles": sc.root_files, "rootKinds": sc.root_kinds,
        "mapping": mapping, "naming": naming, "partial": bool(ctx.get("partial")),
    }


# ───────────────────────────── цикл запросов ─────────────────────────────

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=str(Path.home() / ".hranilka"))
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stdin.reconfigure(encoding="utf-8")
    global THUMBS
    THUMBS = Path(args.data) / "thumbs"
    eng = Engine(Path(args.data))
    emit({"event": "ready"})

    def handle(req: dict):
        rid = req.get("id")
        cmd = req.get("cmd")
        try:
            if cmd == "ping":
                res = {"pong": True}
            elif cmd == "scan":
                res = scan_summary(scan(Path(req["root"]), set(req.get("unskip") or []), set(req.get("exclude") or [])))
            elif cmd == "analyze":
                if eng.job and eng.job.is_alive():
                    eng.stop.set()
                    eng.job.join(timeout=30)
                eng.job = threading.Thread(target=eng.analyze, args=(req,), daemon=True)
                eng.job.start()
                res = {"started": True}
            elif cmd == "pause":
                (eng.paused.set if req.get("on") else eng.paused.clear)()
                res = {"paused": eng.paused.is_set()}
            elif cmd == "stop_and_plan":
                eng.stop.set()
                eng.paused.clear()
                res = {"stopping": True}
            elif cmd == "tune":
                res = eng.tune(req)
            elif cmd == "describe_image":
                res = eng.describe_image(req)
            elif cmd == "learn":
                res = eng.learn(req)
            elif cmd == "replan":
                res = eng.replan(req)
            elif cmd == "learned_info":
                res = eng.learned_info()
            elif cmd == "learned_reset":
                res = eng.learned_reset()
            else:
                raise RuntimeError(f"неизвестная команда {cmd}")
            emit({"id": rid, "ok": True, "result": res})
        except Exception as e:
            traceback.print_exc(file=sys.stderr)
            emit({"id": rid, "ok": False, "error": str(e)})

    for line in sys.stdin:
        line = line.strip()
        if not line:
            continue
        try:
            req = json.loads(line)
        except Exception:
            continue
        threading.Thread(target=handle, args=(req,), daemon=True).start()


if __name__ == "__main__":
    main()
